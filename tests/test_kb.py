"""문서 지식베이스 — 설계 결정을 지키는 테스트.

여기 있는 것은 성능 테스트가 아니다. 누가 나중에 편의를 위해 게이트를 우회하거나
ID 규칙을 바꾸거나 애매한 것을 확정하면 여기서 걸려야 한다. 규제 온톨로지 테스트가
조문 앵커·권한 3분할·삭제 없음을 고정하는 것과 같은 역할이다.
"""

from __future__ import annotations

import json
import re
from pathlib import Path

import pytest

from llmwiki.kb import classify, gate, ingest as kb_ingest, ontology, parse, store, taxonomy
from llmwiki.kb.ingest import UNMASKABLE_RULES, _allowed_after_masking


# --------------------------------------------------------------------------- #
# 택소노미 — 닫힌 집합
# --------------------------------------------------------------------------- #
def test_sector_set_is_closed():
    """미정의 업종이 조용히 통과하면 라벨이 오염된다."""
    with pytest.raises(KeyError):
        taxonomy.get("존재하지_않는_업종")


def test_every_sector_has_required_metrics_and_unit_basis():
    for code, prof in taxonomy.SECTORS.items():
        assert prof.required_metrics, f"{code} 에 필수지표가 없다"
        assert prof.unit_basis, f"{code} 에 원단위 기준이 없다"
        assert prof.name_en and prof.unit_basis_en, f"{code} 에 영어 표기가 없다"


def test_every_required_metric_has_a_label_and_a_pattern():
    """지표를 추가하면서 라벨이나 탐지 어휘를 빠뜨리면 커버리지가 조용히 0이 된다."""
    for prof in taxonomy.SECTORS.values():
        for m in prof.required_metrics:
            assert m in taxonomy.METRIC_LABELS, f"{m} 라벨 없음"
            assert taxonomy.METRIC_PATTERNS.get(m), f"{m} 탐지 어휘 없음"


def test_partition_is_deterministic_and_separates_sectors():
    assert taxonomy.partition("waste") == taxonomy.partition("waste")
    assert taxonomy.partition("waste") != taxonomy.partition("building")


def test_unit_basis_differs_across_sectors():
    """원단위 분모가 같다면 업종을 가를 이유가 없다 — 이 전제가 깨지면 설계가 무의미하다."""
    assert taxonomy.get("building").unit_basis != taxonomy.get("waste").unit_basis


# --------------------------------------------------------------------------- #
# 분류 — 정밀도 우선
# --------------------------------------------------------------------------- #
def test_waste_report_classifies_as_waste():
    text = "음식물류 폐기물 처리시설 퇴비화 부숙 함수율 탈수케이크 자원화 " * 5
    c = classify.classify_text(text)
    assert c.sector == "waste"
    assert not c.needs_review


def test_ambiguous_text_is_held_not_guessed():
    """애매하면 확정하지 않는다. 틀린 라벨은 없는 라벨보다 나쁘다."""
    c = classify.classify_text("보일러 점검 결과 이상 없음")
    assert c.needs_review, "근거가 약한데 업종을 확정했다"


def test_empty_text_falls_back_to_other():
    c = classify.classify_text("")
    assert c.sector == taxonomy.UNCLASSIFIED
    assert c.needs_review
    assert c.method == "fallback"


def test_classification_is_reproducible():
    """같은 입력은 같은 결과. LLM 에게 분류를 물으면 이게 깨진다."""
    text = "음식물 폐기물 퇴비 슬러지 건조기 " * 4
    a, b = classify.classify_text(text), classify.classify_text(text)
    assert (a.sector, round(a.confidence, 6)) == (b.sector, round(b.confidence, 6))
    assert [v.sector for v in a.votes] == [v.sector for v in b.votes]


def test_many_distinct_terms_beat_one_repeated_term():
    """머리말에 반복되는 한 단어로 업종이 뒤집히면 안 된다 (제곱근 가중)."""
    repeated = classify._score("보일러 " * 100)
    diverse = classify._score("음식물 폐기물 퇴비 부숙 슬러지 자원화 함수율 매립 소각 발효")
    assert diverse[0].score > repeated[0].score


def test_manual_override_is_recorded_as_manual():
    """사람이 지정한 것과 룰이 정한 것은 구분되어야 한다."""
    c = classify.manual("building")
    assert (c.sector, c.method, c.needs_review) == ("building", "manual", False)
    with pytest.raises(KeyError):
        classify.manual("없는업종")


def test_metric_coverage_finds_what_is_missing():
    doc = _doc([_table(1, 0, ["구분", "값"], [["연간 전력사용량", "1,200kWh"]])])
    cov = classify.metric_coverage(doc, "waste")
    codes = {m["code"] for m in cov["missing"]}
    assert "annual_ghg_tco2eq" in codes
    assert cov["coverage"] < 1.0
    assert all(m["evidence"] for m in cov["present"]), "근거 없이 present 로 셌다"


# --------------------------------------------------------------------------- #
# 개인정보 — 탐지와 비식별
# --------------------------------------------------------------------------- #
def test_detects_spaced_korean_names():
    """한글 문서는 자간을 벌려 조판한다. '허 만 수' 를 못 잡으면 탐지가 무의미하다."""
    hits = gate.detect_pii("담 당 자 : 허 만 수")
    assert any(h["kind"] == "name" for h in hits)


def test_waste_codes_are_not_account_numbers():
    """51-38-01 은 폐기물 분류코드다. 계좌번호로 잡히면 오탐이 판정을 왜곡한다."""
    hits = gate.detect_pii("음식물류 폐기물 51-38-01 중간가공 51-38-02")
    assert not any(h["kind"] == "account" for h in hits)


def test_business_number_counted_once():
    """사업자등록번호가 계좌번호로도 잡혀 건수가 부풀면 안 된다."""
    hits = gate.detect_pii("사업자등록번호 623-86-00165")
    values = [h["value"] for h in hits]
    assert len(values) == len(set(values))


def test_masking_removes_everything_it_claims_to():
    text = "대표자: 허인구 담당자: 허만수 (063)635-8991 vitech1200@naver.com 623-86-00165"
    v = gate.verify_masking(text)
    assert v["clean"], f"마스킹 후 잔존: {v['residual']}"
    assert v["masked_count"] >= 5


def test_masking_preserves_the_kind():
    """값을 지우지 않고 종류를 남긴다 — '여기 연락처가 있었다' 는 사실이 보존돼야 한다."""
    masked, _ = gate.mask_text("문의 (063)635-8991")
    assert "[전화번호]" in masked
    assert "635-8991" not in masked


# --------------------------------------------------------------------------- #
# 적재 게이트
# --------------------------------------------------------------------------- #
def test_pii_bound_for_overseas_triggers_cross_border_blocker():
    """개인정보가 있는 문서를 국외 모델로 보내면 국외 이전이다 (제28조의8)."""
    r = gate.review("담당자: 허만수 (063)635-8991", destination=gate.destination_for("claude"))
    assert r["verdict"] == "BLOCKED"
    assert not r["upload_allowed"]
    assert any(f["rule"] == "privacy.cross_border" for f in r["findings"])


def test_same_document_is_not_cross_border_on_local_provider():
    """★ 이식하면서 바뀐 판정. LLMWiki 는 사내 모델로 돌릴 수 있으므로 목적지가
    국내면 국외 이전 요건이 적용되지 않는다. 목적지를 하드코딩하면 사내로 돌려도
    차단이 뜨거나, 더 나쁘게는 외부로 보내면서 통과가 뜬다."""
    text = "담당자: 허만수 (063)635-8991"
    local = gate.review(text, destination=gate.destination_for("ollama"))
    assert local["upload_allowed"], "사내 목적지인데 차단됐다"
    assert not any(f["rule"] == "privacy.cross_border" for f in local["findings"])
    assert any(f["rule"] == "privacy.domestic_only" for f in local["findings"])
    # 최소수집 의무는 목적지와 무관하게 남는다
    assert any(f["rule"] == "privacy.minimization" for f in local["findings"])


def test_unknown_provider_is_treated_as_overseas():
    """새 공급자가 추가됐을 때 조용히 통과하는 쪽이 아니라 막히는 쪽으로 틀려야 한다."""
    assert gate.destination_for("someone-new").cross_border
    assert gate.destination_for(None).cross_border


def test_clean_text_allows_upload():
    r = gate.review("루츠블로워 28대의 소비전력을 측정하였다.",
                    destination=gate.destination_for("claude"),
                    has_output_labeling=True, has_prior_notice=True)
    assert r["upload_allowed"]
    assert r["pii_detected"] == 0
    assert r["verdict"] == "ALLOWED"


def test_generative_ai_without_labeling_is_a_violation():
    """AI기본법 제31조제2항 — 생성물 표시."""
    fs = gate.check_ai_act(has_output_labeling=False, has_prior_notice=True)
    assert any(f.rule == "ai.transparency.labeling" and f.severity == "error" for f in fs)


def test_prior_notice_missing_is_a_violation():
    """제31조제1항 — 사전 고지. 화면이 고지를 띄우면 has_prior_notice=True 로 들어온다."""
    fs = gate.check_ai_act(has_output_labeling=True, has_prior_notice=False)
    assert any(f.rule == "ai.transparency.notice" and f.severity == "error" for f in fs)


def test_high_impact_is_held_not_decided():
    """고영향 해당성은 정성 판단이다. 룰이 확정하면 안 된다."""
    hi = next(f for f in gate.check_ai_act() if f.rule == "ai.high_impact.review")
    assert hi.severity == "info"
    assert "판단 유보" in hi.title


def test_findings_have_empty_resolution():
    """판정은 룰이, 확정은 사람이. resolution 을 룰이 채우면 그 경계가 사라진다."""
    r = gate.review("담당자: 허만수 (063)635-8991")
    assert r["findings"], "지적사항이 하나도 없다"
    assert all(f["resolution"] is None for f in r["findings"])


def test_sensitive_data_is_not_resolved_by_masking():
    """주민등록번호는 토큰으로 바꿔도 처리 근거가 생기지 않는다."""
    report = gate.review("담당자 주민등록번호 900101-1234567",
                         destination=gate.destination_for("ollama"))
    masking = gate.verify_masking("담당자 주민등록번호 900101-1234567")
    assert masking["clean"]
    assert "privacy.sensitive" in UNMASKABLE_RULES
    assert not _allowed_after_masking(report, masking), "민감정보인데 마스킹으로 통과됐다"


def test_verdict_labels_cover_every_verdict():
    for lang in ("ko", "en"):
        for v in gate.VERDICTS:
            assert gate.VERDICT_LABELS[lang][v]


# --------------------------------------------------------------------------- #
# 파싱 — 채널과 청킹
# --------------------------------------------------------------------------- #
def test_table_stays_one_chunk():
    """표를 문장처럼 자르면 행-열 관계가 다시 깨진다."""
    doc = _doc([_table(3, 0, ["설비", "용량(kW)"], [["송풍기", "22"], ["건조기", "45"]])])
    chunks = parse.to_chunks(doc)
    tables = [c for c in chunks if c["channel"] == "table"]
    assert len(tables) == 1
    assert "송풍기 | 22" in tables[0]["content"]
    assert "건조기 | 45" in tables[0]["content"]


def test_chunk_channels_are_from_the_closed_set():
    doc = _doc([_table(1, 0, ["a", "b"], [["1kW", "2kW"]])])
    doc.text_blocks.append(parse.TextBlock(page=1, idx=0, text="가" * 40))
    doc.images.append(parse.ImageBlock(page=1, idx=0, width=300, height=200))
    for c in parse.to_chunks(doc):
        assert c["channel"] in parse.CHANNELS


def test_image_kind_comes_from_the_closed_set():
    """종류 추정은 제안이지만, 값 자체는 닫힌 집합이어야 한다."""
    cases = [(20, 20, ""), (400, 300, ""), (400, 300, "배치도"), (400, 300, "추이"),
             (900, 100, "")]
    for w, h, text in cases:
        assert parse._classify_image(w, h, text) in parse.IMAGE_KINDS


def test_logos_do_not_become_search_units():
    doc = _doc([])
    doc.images.append(parse.ImageBlock(page=1, idx=0, width=20, height=20, kind="logo"))
    assert parse.to_chunks(doc) == []


def test_missing_pdfplumber_is_a_clear_error(monkeypatch, tmp_path):
    """설치가 안 됐을 때 ImportError 원문이 아니라 무엇을 어떻게 할지가 나와야 한다."""
    import builtins

    real_import = builtins.__import__

    def fake(name, *args, **kwargs):
        if name == "pdfplumber":
            raise ImportError("no pdfplumber")
        return real_import(name, *args, **kwargs)

    monkeypatch.setattr(builtins, "__import__", fake)
    pdf = tmp_path / "x.pdf"
    pdf.write_bytes(b"%PDF-1.4")
    with pytest.raises(parse.ParseError, match="pdfplumber"):
        parse.parse_pdf(str(pdf))


# --------------------------------------------------------------------------- #
# 온톨로지 — ID 규칙이 단일 실패 지점
# --------------------------------------------------------------------------- #
def test_node_ids_contain_no_coordinates():
    """★ 단일 실패 지점. ID 에 bbox 를 넣으면 파서를 고치는 순간 그래프가 끊긴다."""
    g = _graph()
    for n in g["nodes"]:
        assert not re.search(r"\bbbox\b|\bx0\b|\d+\.\d{3,}", n["id"]), \
            f"ID 에 좌표로 보이는 값이 있다: {n['id']}"


def test_validator_rejects_a_coordinate_in_an_id():
    """검증기가 실제로 잡는지 확인한다 — 안 잡으면 위 테스트는 장식이다."""
    g = _graph()
    g["nodes"][0]["id"] = "dgn:t1/bbox_72.0_531.44"
    assert not ontology.validate_graph(g).ok


def test_graph_is_byte_identical_on_rerun():
    assert json.dumps(_graph(), sort_keys=True) == json.dumps(_graph(), sort_keys=True)


def test_header_unit_is_borrowed_for_bare_numbers():
    """`용량(kW)` 헤더 아래 `22` 는 실제 진단서에서 가장 흔한 표기다."""
    qs = [n for n in _graph()["nodes"] if n["type"] == "Quantity"]
    assert any(q["unit"] == "kW" and q["value"] == 22.0 for q in qs), \
        f"헤더 단위를 빌려오지 못했다: {[(q['value'], q['unit']) for q in qs]}"


def test_korean_word_is_not_read_as_currency():
    """'남원시' 의 '원' 을 화폐 단위로 읽으면 안 된다."""
    doc = _doc([_table(1, 0, ["소재지"], [["전라북도 남원시 대강면 섬진로 1200-27"]])])
    g = ontology.build_graph(doc, classify.manual("waste"), {"missing": [], "present": []})
    assert not [n for n in g["nodes"] if n["type"] == "Quantity" and n["unit"] == "원"]


def test_year_is_not_counted_as_a_quantity_dimension():
    """'2015년' 을 기간으로 읽으면 에너지 집계가 오염된다."""
    doc = _doc([_table(1, 0, ["구분", "값"], [["설립년도", "2015년"]])])
    g = ontology.build_graph(doc, classify.manual("waste"), {"missing": [], "present": []})
    years = [n for n in g["nodes"] if n["type"] == "Quantity" and n["unit"] == "년"]
    assert years and all(y["dimension"] in ontology.NON_QUANTITY_DIMENSIONS for y in years)


def test_every_quantity_has_evidence():
    """근거 없는 사실 금지."""
    g = _graph()
    qids = {n["id"] for n in g["nodes"] if n["type"] == "Quantity"}
    evidenced = {e["source"] for e in g["edges"] if e["type"] == "evidencedBy"}
    assert qids and qids <= evidenced, f"근거 없는 Quantity: {qids - evidenced}"


def test_validator_rejects_a_quantity_without_evidence():
    g = _graph()
    g["edges"] = [e for e in g["edges"] if e["type"] != "evidencedBy"]
    result = ontology.validate_graph(g)
    assert not result.ok
    assert any(i.code == "node.span" for i in result.errors)


def test_validator_rejects_a_rule_filled_resolution():
    """룰이 resolution 을 채우면 판정과 확정의 구분이 사라진다."""
    doc = _doc([_table(1, 0, ["구분", "용량(kW)"], [["송풍기", "22"]])])
    report = {"findings": [{
        "rule": "privacy.minimization", "severity": "error", "law": "개인정보보호법",
        "article": "제3조", "title": "t", "detail": "d", "resolution": "룰이 해결함",
    }]}
    g = ontology.build_graph(doc, classify.manual("waste"),
                            {"missing": [], "present": []}, report)
    assert any(i.code == "finding.resolution" for i in ontology.validate_graph(g).errors)


def test_every_node_has_a_derivation_from_the_closed_set():
    for n in _graph()["nodes"]:
        assert n.get("derivation") in ontology.DERIVATIONS, \
            f"{n['id']} 의 derivation 이 허용값 {ontology.DERIVATIONS} 밖이다"


def test_built_graph_passes_its_own_schema():
    result = ontology.validate_graph(_graph())
    assert result.ok, [str(i) for i in result.errors]


def test_schema_dict_declares_every_type():
    s = ontology.schema_dict()
    assert set(s["nodes"]) == set(ontology.NODE_TYPES)
    assert set(s["edges"]) == set(ontology.EDGE_TYPES)
    assert s["ontology"] == ontology.KB_ONTOLOGY_VERSION


def test_turtle_export_is_parseable_shape():
    ttl = ontology.to_turtle(_graph())
    assert ttl.startswith("@prefix ed:")
    assert " rdf:type ed:Diagnosis ;" in ttl


def test_real_diagnosis_graph_still_validates():
    """실측 산출물(폐기물처리 진단서 32면, 노드 101개)로 스키마를 고정한다.

    원 구현(RAG-AI_Gov)이 낸 그래프가 이식본의 온톨로지를 그대로 통과해야 한다 —
    통과하지 못하면 이식 과정에서 형식이 갈라진 것이다. 픽스처는 고객사 식별자를
    지운 것이고(구조는 원본 그대로), 개인정보는 애초에 그래프에 올라오지 않는다.
    """
    fixture = Path(__file__).parent / "data" / "diagnosis_graph.json"
    graph = json.loads(fixture.read_text(encoding="utf-8"))
    result = ontology.validate_graph(graph)
    assert result.ok, [str(i) for i in result.errors]
    assert graph["stats"]["nodes"] == len(graph["nodes"])


# --------------------------------------------------------------------------- #
# 저장소 — 게이트에 우회로가 없어야 한다
# --------------------------------------------------------------------------- #
class _Result:
    """저장소가 보는 최소한의 분석 결과."""

    def __init__(self, *, allowed: bool = True, needs_review: bool = False):
        self.upload_allowed = allowed
        self.needs_review = needs_review
        self.partition = taxonomy.partition("waste")
        self.filename = "x.pdf"
        self.doc_hash = "abc123"
        self.sector = "waste"
        self.sector_name = "폐기물처리·자원순환"
        self.gate = {"pii_detected": 0, "verdict": "ALLOWED"}
        self.graph_stats = {"nodes": 3}
        self.graph = None
        self.excel_path = None
        self.chunks = [{"channel": "text", "content": "본문 " * 20, "page": 1, "anchor": "p1/t0"}]

    def to_dict(self):
        return {"doc_hash": self.doc_hash, "sector": self.sector}


def test_store_refuses_when_the_gate_said_no(tmp_path):
    """게이트를 우회하는 인자는 없다 — allowed=False 면 아무것도 쓰이지 않아야 한다."""
    s = store.Store(tmp_path)
    out = s.ingest(_Result(allowed=False))
    assert out["stored"] == 0
    assert "규제 게이트" in out["skipped"]
    assert s.documents() == []
    assert not s.ledger_path.exists()


def test_store_refuses_an_unconfirmed_sector(tmp_path):
    """잘못 분류된 문서는 영영 엉뚱한 구획에서 검색된다."""
    out = store.Store(tmp_path).ingest(_Result(needs_review=True))
    assert out["stored"] == 0
    assert "업종" in out["skipped"]


def test_prepare_masks_and_verifies():
    chunks = [{"channel": "text", "content": "담당자 홍길동, 010-1234-5678"}]
    cleaned, info = store.prepare(chunks)
    assert info["masked"] is True
    assert info["residual_count"] == 0
    assert "010-1234-5678" not in cleaned[0]["content"]


def test_prepare_refuses_unmasked_text_that_still_has_pii():
    """mask=False 로 우회하려 해도 개인정보가 남아 있으면 빈 목록을 준다."""
    chunks = [{"channel": "text", "content": "연락처 010-1234-5678"}]
    cleaned, info = store.prepare(chunks, mask=False)
    assert cleaned == []
    assert info["residual_count"] > 0


def test_store_stops_when_masking_leaves_something(tmp_path):
    s = store.Store(tmp_path)
    r = _Result()
    r.chunks = [{"channel": "text", "content": "연락처 010-1234-5678", "page": 1, "anchor": "a"}]
    out = s.ingest(r, mask=False)
    assert out["stored"] == 0
    assert "개인정보" in out["skipped"]
    assert s.documents() == []


def test_stored_channels_carry_no_pii(tmp_path):
    s = store.Store(tmp_path)
    r = _Result()
    r.chunks = [{"channel": "text", "content": "담당자: 허만수 (063)635-8991 " * 3,
                 "page": 1, "anchor": "p1/t0"}]
    assert s.ingest(r)["stored"] == 1
    for rec in s.channels(r.doc_hash):
        assert not gate.detect_pii(rec["content"]), "적재본에 개인정보가 남았다"
        assert "[전화번호]" in rec["content"], "종류 토큰까지 지워졌다"


def test_each_table_becomes_its_own_search_unit():
    """표 두 개를 한 레코드에 넣으면 검색이 엉뚱한 표를 근거로 답한다."""
    chunks = [
        {"channel": "table", "content": "표1", "page": 1, "anchor": "t1"},
        {"channel": "table", "content": "표2", "page": 2, "anchor": "t2"},
    ]
    docs = store.channel_documents(chunks)
    assert len(docs) == 2
    assert all(d["parts"] == 1 for d in docs)


def test_text_chunks_are_merged_but_bounded():
    chunks = [{"channel": "text", "content": "가" * 500, "page": i, "anchor": f"p{i}"}
              for i in range(20)]
    docs = store.channel_documents(chunks, max_chars=2000)
    assert len(docs) > 1
    assert all(len(d["content"]) < 4000 for d in docs)


def test_ledger_is_append_only(tmp_path):
    """적재 이력이 덮어써지면 어떤 문서가 언제 들어왔는지 감사에 답할 수 없다."""
    s = store.Store(tmp_path)
    s.ingest(_Result())
    s.ingest(_Result())
    assert len(s.ledger()) == 2
    assert len(s.documents()) == 1, "같은 문서는 목록에 한 번만 나와야 한다"


def test_search_filters_by_sector_and_channel(tmp_path):
    s = store.Store(tmp_path)
    r = _Result()
    r.chunks = [
        {"channel": "text", "content": "송풍기 소비전력 측정 " * 5, "page": 1, "anchor": "p1/t0"},
        {"channel": "table", "content": "송풍기 | 22kW", "page": 2, "anchor": "p2/tbl0"},
    ]
    s.ingest(r)
    assert s.search("송풍기")
    assert all(h["channel"] == "table" for h in s.search("송풍기", channel="table"))
    assert s.search("송풍기", sector="building") == [], "업종 필터가 실제로 걸리지 않았다"


def test_search_needs_every_term(tmp_path):
    """한 낱말만 걸려도 올라오면 수치 질의의 근거가 흐려진다 — 모든 낱말이 있어야 한다."""
    s = store.Store(tmp_path)
    r = _Result()
    r.chunks = [{"channel": "table", "content": "송풍기 | 22kW | 7,200h",
                 "page": 2, "anchor": "p2/tbl0"}]
    s.ingest(r)
    assert s.search("송풍기 22kW")
    assert s.search("송풍기 건조기") == []


# --------------------------------------------------------------------------- #
# 도우미 — PDF 없이 파싱 결과를 흉내 낸다
# --------------------------------------------------------------------------- #
def _table(page: int, idx: int, header: list[str], rows: list[list[str]]) -> parse.TableBlock:
    return parse.TableBlock(page=page, idx=idx, header=header, rows=rows)


def _doc(tables: list[parse.TableBlock]) -> parse.ParsedDocument:
    return parse.ParsedDocument(
        filename="t.pdf", doc_hash="deadbeefdeadbeef", n_pages=1, tables=tables,
    )


def _graph() -> dict:
    doc = _doc([_table(7, 0, ["구분", "용량(kW)", "수량"], [["루츠블로워", "22", "18대"]])])
    doc.text_blocks.append(parse.TextBlock(page=7, idx=0, text="루츠블로워 보일러 점검"))
    return ontology.build_graph(
        doc, classify.manual("waste"), {"missing": [], "present": []}, diagnosis_id="t1",
    )


# --------------------------------------------------------------------------- #
# 목적지 선택 · 채널별 내용
# --------------------------------------------------------------------------- #
def test_selectable_destinations_are_a_closed_set():
    """화면 드롭다운의 유일한 출처. 국외 이전 해당성을 서버가 함께 내려준다 —
    화면이 이름만 보고 국내/국외를 추측하면 두 곳이 어긋난다."""
    out = gate.selectable_destinations()
    assert [d["provider"] for d in out] == list(gate.SELECTABLE_PROVIDERS)
    by = {d["provider"]: d for d in out}
    assert by["ollama"]["cross_border"] is False
    assert by["grok"]["cross_border"] is True
    assert all(d["name"] for d in out)


def test_the_same_document_is_judged_differently_by_destination():
    """사내 GPU 로 돌리면 국외 이전 조항이 걸리지 않는다. 같은 문서·같은 룰이라도
    어디로 보내느냐가 판정을 가른다 — 이게 공급자 선택의 존재 이유다."""
    text = "담당자 홍길동 010-1234-5678 이 확인한 루츠블로워 22kW 18대"
    rules = lambda provider: {                                       # noqa: E731
        f["rule"] for f in gate.review(
            text, destination=gate.destination_for(provider), masking_enabled=False,
        )["findings"]
    }
    assert "privacy.cross_border" in rules("grok")
    assert "privacy.cross_border" not in rules("ollama")


def test_preview_matches_the_channel_counts():
    """'채널' 카드의 개수와 목록의 줄 수가 어긋나면 어느 쪽이 맞는지 알 수 없다.
    글은 짧은 조각을, 그림은 로고를 빼는 기준까지 `to_chunks` 와 같아야 한다."""
    doc = _doc([_table(7, 0, ["구분", "용량(kW)"], [["루츠블로워", "22"]])])
    doc.text_blocks.append(parse.TextBlock(page=7, idx=0, text="쪽번호"))          # 20자 미만
    doc.text_blocks.append(parse.TextBlock(page=7, idx=1, text="루츠블로워 배열이 회수되지 않아 손실이 크다"))
    doc.images.append(parse.ImageBlock(page=7, idx=0, width=100, height=80, kind="photo"))
    doc.images.append(parse.ImageBlock(page=7, idx=1, width=40, height=40, kind="logo"))

    preview = kb_ingest._build_preview(doc, mask=False)
    chunks = parse.to_chunks(doc)

    for channel in ("text", "table"):
        assert len(preview[channel]) == sum(1 for c in chunks if c["channel"] == channel)
    # 로고는 목록에 남기되 적재 대상에서는 빠진다 — 조용히 사라지면 개수 차이를 설명할 수 없다
    assert len(preview["image"]) == 2
    assert [i["indexed"] for i in preview["image"]] == [True, False]
    assert sum(1 for c in chunks if c["channel"] == "image") == 1


def test_preview_keeps_the_table_grid():
    """표를 문장으로 이어 붙이면 행-열 관계가 다시 깨진다. 화면이 격자로 그릴 수
    있도록 머리행과 행을 나눠서 준다."""
    doc = _doc([_table(7, 0, ["구분", "용량(kW)"], [["루츠블로워", "22"], ["보일러", "0.5"]])])
    tables = kb_ingest._build_preview(doc, mask=False)["table"]
    assert len(tables) == 1
    assert tables[0]["header"] == ["구분", "용량(kW)"]
    assert tables[0]["rows"] == [["루츠블로워", "22"], ["보일러", "0.5"]]
    assert tables[0]["anchor"] == "p7/tbl0"


def test_preview_hides_what_the_gate_blocked():
    """게이트가 원문 적재를 막은 문서를 같은 화면이 원문 그대로 보여주면 막은 의미가 없다."""
    doc = _doc([_table(7, 0, ["담당"], [["홍길동 010-1234-5678"]])])
    doc.text_blocks.append(parse.TextBlock(page=7, idx=0, text="담당자 홍길동 010-1234-5678 확인함"))

    masked = kb_ingest._build_preview(doc, mask=True)
    assert masked["masked"] is True
    assert "010-1234-5678" not in masked["text"][0]["content"]
    assert "010-1234-5678" not in masked["table"][0]["rows"][0][0]

    raw = kb_ingest._build_preview(doc, mask=False)
    assert raw["masked"] is False
    assert "010-1234-5678" in raw["text"][0]["content"]


# --------------------------------------------------------------------------- #
# 입력 형식 — PDF · 엑셀 · 이미지가 같은 ParsedDocument 로 들어온다
# --------------------------------------------------------------------------- #
def _workbook(path, sheets: dict[str, list[list]]):
    from openpyxl import Workbook

    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(title=name)
        for row in rows:
            ws.append(row)
    wb.save(path)
    return str(path)


def test_unknown_suffix_is_rejected_not_guessed():
    """확장자를 짐작하면 엑셀이 PDF 파서로 들어가 스택트레이스가 난다."""
    from llmwiki.kb import sources

    with pytest.raises(parse.ParseError):
        sources.kind_of("보고서.docx")
    assert sources.kind_of("계측.XLSX") == "sheet"
    assert sources.kind_of("명판.JPG") == "image"


def test_spreadsheet_splits_blocks_on_blank_rows(tmp_path):
    """한 시트에 표가 둘이면 둘로 읽어야 한다. 통째로 묶으면 헤더가 엉뚱한 줄이 된다."""
    from llmwiki.kb import sources

    path = _workbook(tmp_path / "계측.xlsx", {
        "측정": [
            ["기번", "측정전력(kW)", "부하율(%)"],
            ["#1", 25.7, 117],
            ["#3", 23.0, 105],
            [],
            ["구분", "연간 전력량(kWh/y)"],
            ["2차 숙성실", 2664576],
        ],
    })
    doc = sources.parse_spreadsheet(path)
    assert doc.n_pages == 1
    assert len(doc.tables) == 2
    assert doc.tables[0].header[0] == "기번"
    assert doc.tables[1].header[1] == "연간 전력량(kWh/y)"
    assert doc.tables[0].caption == "측정"
    assert doc.n_numeric_cells >= 5


def test_spreadsheet_drops_layout_padding_columns(tmp_path):
    """서식용 빈 열이 남으면 헤더가 `['', '', '항목']` 이 되어 열 이름 탐색이 헛돈다."""
    from llmwiki.kb import sources

    path = _workbook(tmp_path / "여백.xlsx", {
        "표": [["", "", "항목", "값"], ["", "", "용량", "22kW"]],
    })
    doc = sources.parse_spreadsheet(path)
    assert doc.tables[0].header == ["항목", "값"]


def test_spreadsheet_feeds_the_same_pipeline_as_pdf(tmp_path):
    """업종 분류·게이트·온톨로지는 형식을 모른다 — 같은 ParsedDocument 만 본다."""
    from llmwiki.kb import sources

    path = _workbook(tmp_path / "폐기물.xlsx", {
        "진단": [
            ["항목", "값"],
            ["음식물 폐기물 처리량", "48톤/일"],
            ["함수율", "80%"],
            ["루츠블로워 소비전력", "25.7kW"],
        ],
    })
    doc = sources.parse_document(path)
    cls = classify.classify_document(doc)
    assert cls.sector == "waste"
    assert parse.to_chunks(doc), "채널 청크가 나와야 적재된다"


def test_broken_file_says_what_failed(tmp_path):
    """확장자는 맞는데 내용이 깨진 경우 — 스택트레이스가 아니라 문장을 돌려준다."""
    from llmwiki.kb import sources

    path = tmp_path / "깨진.xlsx"
    path.write_bytes(b"not a workbook")
    with pytest.raises(parse.ParseError) as exc:
        sources.parse_document(str(path))
    assert "깨진.xlsx" in str(exc.value)


def test_image_without_ocr_says_it_could_not_read(tmp_path):
    """빈 결과를 주면 '아무것도 없는 사진' 과 '읽지 못했다' 가 구분되지 않는다."""
    from PIL import Image

    from llmwiki.kb import sources

    path = tmp_path / "명판.png"
    Image.new("RGB", (400, 300), "white").save(path)
    doc = sources.parse_document(str(path))

    assert doc.n_pages == 1
    assert len(doc.images) == 1 and doc.images[0].width == 400
    if not sources.ocr_ready()["ok"]:
        assert any("읽지 못했다" in w for w in doc.warnings)


def test_readiness_reports_each_format_separately():
    """이미지 OCR 만 빠진 상태는 정상이다. 전체를 실패로 칠하면 PDF 경로까지 막힌 것처럼 보인다."""
    from llmwiki.kb import sources

    r = sources.readiness()
    assert r["pdf"]["ok"] and r["sheet"]["ok"]
    assert set(r["image"]) >= {"ok", "reason", "hint"}
    assert ".xlsx" in r["suffixes"] and ".png" in r["suffixes"]
